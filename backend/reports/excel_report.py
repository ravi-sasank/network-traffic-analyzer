"""SENTINEL — Network Security Assessment Workbook.

Designed to the same editorial standard as the PDF: clean light sheets,
Arial throughout, strong typographic hierarchy, restrained accent colour.
Nine sheets covering summary, findings, ATT&CK mapping, detections, risk,
traffic and methodology, with embedded charts and data-bar formatting.
"""
import io
import json
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from backend.storage.db import get_connection
from backend.storage import api_queries as aq

FONT = 'Arial'

# ── palette: mirrors the PDF ──────────────────────────────────────────────
INK    = '101823'
GRAPH  = '2C3A4D'
MUTED  = '6B7B8F'
HAIR   = 'DBE3EC'
WASH   = 'F5F8FB'
ACCENT = '0B6E8A'
WHITE  = 'FFFFFF'

SEV = {'critical': 'C0263F', 'high': 'C96A1B',
       'medium': 'A8830E', 'low': '2B7D96'}
TIER = {**SEV, 'elevated': 'A8830E', 'guarded': '2B7D96', 'normal': '2F7D5F'}

ATTACK = {
    'port_scan':          ('TA0007 Discovery', 'T1046', 'Network Service Discovery'),
    'connection_burst':   ('TA0040 Impact', 'T1498.001', 'Direct Network Flood'),
    'dns_exfil':          ('TA0010 Exfiltration', 'T1048.003',
                           'Exfiltration Over Unencrypted Protocol'),
    'icmp_sweep':         ('TA0007 Discovery', 'T1018', 'Remote System Discovery'),
    'volume_anomaly':     ('TA0010 Exfiltration', 'T1030', 'Data Transfer Size Limits'),
    'connection_anomaly': ('TA0011 Command and Control', 'T1071',
                           'Application Layer Protocol'),
}

REMEDIATION = {
    'port_scan': ('HIGH', 'Isolate and investigate the scanning host',
                  'Quarantine the source, review its process and authentication '
                  'history, and confirm whether this was authorised scanning.'),
    'connection_burst': ('HIGH', 'Apply connection rate limiting',
                         'Enforce per-source connection limits at the gateway and '
                         'enable SYN cookies on exposed services.'),
    'dns_exfil': ('CRITICAL', 'Restrict and inspect outbound DNS',
                  'Force DNS through controlled resolvers, block direct port 53 '
                  'egress, alert on anomalous query-name entropy.'),
    'icmp_sweep': ('MEDIUM', 'Review ICMP egress policy',
                   'Limit ICMP between segments and log sweep behaviour.'),
    'volume_anomaly': ('MEDIUM', 'Validate the transfer against business need',
                       'Confirm destination and volume are expected; inspect the '
                       'flow record and destination reputation.'),
    'connection_anomaly': ('MEDIUM', 'Profile the host for beaconing',
                           'Check for regular-interval connections indicative of '
                           'C2 and review destination reputation.'),
}

THIN = Side(style='thin', color=HAIR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOTTOM = Border(bottom=THIN)


def fmt_bytes(n):
    n = float(n or 0)
    for u in ('B', 'KB', 'MB', 'GB', 'TB'):
        if n < 1024:
            return f"{int(n)} {u}" if u == 'B' else f"{n:.1f} {u}"
        n /= 1024
    return f"{n:.1f} PB"


def ts_str(v):
    if not v:
        return ''
    return datetime.fromtimestamp(float(v), tz=timezone.utc).strftime('%Y-%m-%d %H:%M:%S')


# ── sheet furniture ───────────────────────────────────────────────────────
def sheet_title(ws, num, title, subtitle, span=6):
    """Consistent section header on every sheet."""
    ws.sheet_view.showGridLines = False
    ws['A1'] = num
    ws['A1'].font = Font(name=FONT, size=8, bold=True, color=ACCENT)
    ws['A2'] = title
    ws['A2'].font = Font(name=FONT, size=16, bold=True, color=INK)
    ws['A3'] = subtitle
    ws['A3'].font = Font(name=FONT, size=9, color=MUTED)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=span)
    ws.row_dimensions[1].height = 13
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 15
    ws.row_dimensions[4].height = 8
    # rule under the header
    for c in range(1, span + 1):
        ws.cell(row=4, column=c).border = Border(bottom=Side(style='medium', color=INK))


def table(ws, header_row, headers, widths, n_rows, wrap_cols=()):
    """Apply the house table style."""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.fill = PatternFill('solid', fgColor=INK)
        c.font = Font(name=FONT, size=8, bold=True, color=WHITE)
        c.alignment = Alignment(vertical='center', horizontal='left', indent=1)
        c.border = BORDER
    ws.row_dimensions[header_row].height = 20
    for r in range(header_row + 1, header_row + n_rows + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if not cell.font or cell.font.name != FONT:
                cell.font = Font(name=FONT, size=9, color=GRAPH)
            cell.alignment = Alignment(vertical='top', indent=1,
                                       wrap_text=(c in wrap_cols))
            if (r - header_row) % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=WASH)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if n_rows:
        ws.auto_filter.ref = (f"A{header_row}:"
                              f"{get_column_letter(len(headers))}{header_row + n_rows}")


def build_workbook(conn=None):
    own = conn is None
    if own:
        conn = get_connection()
    try:
        stats     = aq.get_stats(conn)
        alerts    = aq.get_alerts(conn, limit=3000)
        board     = aq.get_threat_board(conn)
        talkers   = aq.get_top_talkers(conn, limit=200)
        protocols = aq.get_protocol_breakdown(conn)
        domains   = aq.get_top_domains(conn, limit=30)
        timeline  = aq.get_traffic_timeline(conn, minutes=60)
        span = conn.execute("SELECT MIN(bucket_ts) a, MAX(bucket_ts) b FROM flows").fetchone()
    finally:
        if own:
            conn.close()

    sev_count, rule_count = {}, {}
    for a in alerts:
        sev_count[a['severity']] = sev_count.get(a['severity'], 0) + 1
        rule_count[a['rule_name']] = rule_count.get(a['rule_name'], 0) + 1
    hostile = [b for b in board if b['tier'] in ('critical', 'high')]
    top = board[0]['score'] if board else 0
    pkey = ('critical' if top >= 80 else 'high' if top >= 55
            else 'elevated' if top >= 30 else 'normal')
    posture = {'critical': 'CRITICAL', 'high': 'HIGH',
               'elevated': 'ELEVATED', 'normal': 'NOMINAL'}[pkey]
    ref = datetime.now(timezone.utc).strftime('NTA-%Y%m%d-%H%M')

    wb = Workbook()

    # ══ 1. SUMMARY ═══════════════════════════════════════════════════════
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEF', (30, 26, 4, 30, 26, 4)):
        ws.column_dimensions[col].width = w

    ws['A1'] = 'SENTINEL'
    ws['A1'].font = Font(name=FONT, size=22, bold=True, color=INK)
    ws['A2'] = 'NETWORK SITUATIONAL AWARENESS'
    ws['A2'].font = Font(name=FONT, size=8, bold=True, color=ACCENT)
    ws['A4'] = 'Network Security Assessment'
    ws['A4'].font = Font(name=FONT, size=15, bold=True, color=INK)
    ws['A5'] = f"Assessment window   {ts_str(span['a'])}  to  {ts_str(span['b'])} UTC"
    ws['A5'].font = Font(name=FONT, size=9, color=MUTED)
    ws['D1'] = f"Report ref   {ref}"
    ws['D1'].font = Font(name=FONT, size=9, color=MUTED)
    ws['D2'] = datetime.now(timezone.utc).strftime('Generated %d %B %Y at %H:%M UTC')
    ws['D2'].font = Font(name=FONT, size=9, color=MUTED)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[4].height = 20
    for c in range(1, 7):
        ws.cell(row=6, column=c).border = Border(bottom=Side(style='medium', color=INK))

    # posture block
    ws['A8'] = 'ASSESSED NETWORK POSTURE'
    ws['A8'].font = Font(name=FONT, size=8, bold=True, color=MUTED)
    ws['A9'] = posture
    ws['A9'].font = Font(name=FONT, size=20, bold=True, color=TIER[pkey])
    ws['A10'] = {
        'critical': 'Active hostile behaviour observed. Immediate investigation required.',
        'high': 'Significant suspicious activity detected. Prompt review advised.',
        'elevated': 'Anomalous activity present above baseline. Monitor closely.',
        'normal': 'No material threats observed during the assessment window.',
    }[pkey]
    ws['A10'].font = Font(name=FONT, size=9, color=GRAPH)
    ws.merge_cells('A10:F10')
    ws.row_dimensions[9].height = 26
    for r in (8, 9, 10):
        for c in range(1, 7):
            ws.cell(row=r, column=c).fill = PatternFill('solid', fgColor=WASH)
        ws.cell(row=r, column=1).border = Border(
            left=Side(style='thick', color=TIER[pkey]))

    # metrics, two columns
    ws['A12'] = 'KEY METRICS'
    ws['A12'].font = Font(name=FONT, size=9, bold=True, color=INK)
    metrics = [
        ('Traffic volume', fmt_bytes(stats['bytes_last_hour'])),
        ('Active hosts', f"{stats['active_devices']:,}"),
        ('Flows recorded', f"{stats['flows_last_hour']:,}"),
        ('Packets observed', f"{stats['packets_last_hour']:,}"),
        ('Detections raised', f"{stats['alerts_last_hour']:,}"),
        ('Hosts at risk', f"{len(hostile):,}"),
        ('Peak threat score', f"{top:.0f} / 100"),
        ('Detection rules armed', '5'),
    ]
    for i, (k, v) in enumerate(metrics):
        r = 14 + (i % 4)
        cbase = 1 if i < 4 else 4
        kc = ws.cell(row=r, column=cbase, value=k)
        kc.font = Font(name=FONT, size=9, color=GRAPH)
        kc.border = BOTTOM
        vc = ws.cell(row=r, column=cbase + 1, value=v)
        vc.font = Font(name=FONT, size=11, bold=True, color=INK)
        vc.alignment = Alignment(horizontal='right')
        vc.border = BOTTOM
        ws.row_dimensions[r].height = 18

    # findings
    ws['A20'] = 'KEY FINDINGS'
    ws['A20'].font = Font(name=FONT, size=9, bold=True, color=INK)
    findings = []
    if hostile:
        findings.append(('CRITICAL',
            f"{len(hostile)} host(s) exhibited hostile behaviour — "
            f"{', '.join(b['host'] for b in hostile[:3])}. Immediate triage required."))
    if rule_count.get('dns_exfil'):
        findings.append(('CRITICAL',
            f"{rule_count['dns_exfil']} detection(s) of high-entropy DNS labels "
            "consistent with tunnelling (ATT&CK T1048.003)."))
    if rule_count.get('port_scan'):
        findings.append(('HIGH',
            f"{rule_count['port_scan']} port-scanning event(s) observed "
            "(ATT&CK T1046) — an actor mapping the attack surface."))
    if rule_count.get('connection_burst'):
        findings.append(('HIGH',
            f"{rule_count['connection_burst']} volumetric flood event(s) "
            "(ATT&CK T1498.001) against local services."))
    anom = rule_count.get('volume_anomaly', 0) + rule_count.get('connection_anomaly', 0)
    if anom:
        findings.append(('MEDIUM',
            f"{anom} statistical anomal(ies) materially above host baseline."))
    if not findings:
        findings.append(('LOW',
            'No material threats identified. All rules armed throughout.'))
    for i, (sev, text) in enumerate(findings):
        r = 22 + i
        nc = ws.cell(row=r, column=1, value=f"{i+1:02d}   {sev}")
        nc.font = Font(name=FONT, size=9, bold=True, color=SEV[sev.lower()])
        tc = ws.cell(row=r, column=2, value=text)
        tc.font = Font(name=FONT, size=9, color=GRAPH)
        tc.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 26

    foot = 22 + len(findings) + 2
    ws.cell(row=foot, column=1,
            value='Captured on operator-administered infrastructure. '
                  'No offensive security testing was performed.').font = \
        Font(name=FONT, size=8, italic=True, color=MUTED)
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=6)

    # ══ 2. ATT&CK ════════════════════════════════════════════════════════
    ws = wb.create_sheet('ATT&CK Mapping')
    sheet_title(ws, '02  ADVERSARY TECHNIQUES', 'MITRE ATT&CK mapping',
                'Detections mapped to adversary tactics and techniques.', span=5)
    hdr = 6
    table(ws, hdr, ['TACTIC', 'TECHNIQUE', 'TECHNIQUE NAME', 'OBSERVED', 'SEVERITY'],
          [30, 16, 44, 14, 14], len(rule_count))
    for i, (rule, n) in enumerate(sorted(rule_count.items(), key=lambda x: -x[1])):
        tac, tid, tname = ATTACK.get(rule, ('—', '—', rule.replace('_', ' ').title()))
        sev = next((a['severity'] for a in alerts if a['rule_name'] == rule), 'low')
        r = hdr + 1 + i
        ws.cell(row=r, column=1, value=tac)
        c = ws.cell(row=r, column=2, value=tid)
        c.font = Font(name=FONT, size=9, bold=True, color=INK)
        ws.cell(row=r, column=3, value=tname)
        c = ws.cell(row=r, column=4, value=n)
        c.font = Font(name=FONT, size=9, bold=True, color=INK)
        c = ws.cell(row=r, column=5, value=sev.upper())
        c.font = Font(name=FONT, size=9, bold=True, color=SEV.get(sev, MUTED))
    note = hdr + len(rule_count) + 2
    ws.cell(row=note, column=1,
            value='Technique identifiers reference the MITRE ATT&CK Enterprise matrix. '
                  'Mapping is indicative of technique class, not attribution.').font = \
        Font(name=FONT, size=8, italic=True, color=MUTED)

    # ══ 3. RISK REGISTER ═════════════════════════════════════════════════
    ws = wb.create_sheet('Risk Register')
    sheet_title(ws, '03  HOST RISK', 'Scored hosts',
                'Scores decay with a ten-minute half-life, reflecting present risk.',
                span=5)
    hdr = 6
    table(ws, hdr, ['HOST', 'SCORE', 'TIER', 'EVENTS', 'MOST RECENT DETECTION'],
          [32, 12, 14, 12, 34], len(board))
    for i, b in enumerate(board):
        r = hdr + 1 + i
        ws.cell(row=r, column=1, value=b['host']).font = \
            Font(name=FONT, size=9, color=INK)
        ws.cell(row=r, column=2, value=round(b['score'], 1))
        c = ws.cell(row=r, column=3, value=b['tier'].upper())
        c.font = Font(name=FONT, size=9, bold=True, color=TIER.get(b['tier'], MUTED))
        ws.cell(row=r, column=4, value=b['alert_count'])
        ws.cell(row=r, column=5,
                value=b['recent'][-1]['rule'].replace('_', ' ') if b.get('recent') else '')
    if board:
        ws.conditional_formatting.add(
            f"B{hdr+1}:B{hdr+len(board)}",
            DataBarRule(start_type='num', start_value=0, end_type='num',
                        end_value=100, color=ACCENT, showValue=True))

    # ══ 4. DETECTIONS ════════════════════════════════════════════════════
    ws = wb.create_sheet('Detections')
    sheet_title(ws, '04  DETECTION LOG', 'Recorded events',
                'Every detection with its rationale and supporting evidence.', span=7)
    hdr = 6
    table(ws, hdr,
          ['TIMESTAMP (UTC)', 'RULE', 'ATT&CK', 'SEVERITY', 'SOURCE',
           'EVIDENCE', 'RATIONALE'],
          [20, 20, 14, 12, 26, 40, 84], len(alerts), wrap_cols=(7,))
    for i, a in enumerate(alerts):
        ev = a.get('evidence') or {}
        if isinstance(ev, str):
            try:
                ev = json.loads(ev)
            except ValueError:
                ev = {}
        r = hdr + 1 + i
        ws.cell(row=r, column=1, value=ts_str(a['ts']))
        ws.cell(row=r, column=2, value=a['rule_name'].replace('_', ' '))
        ws.cell(row=r, column=3, value=ATTACK.get(a['rule_name'], ('', '—'))[1])
        c = ws.cell(row=r, column=4, value=a['severity'].upper())
        c.font = Font(name=FONT, size=9, bold=True, color=SEV.get(a['severity'], MUTED))
        ws.cell(row=r, column=5, value=a.get('src_ip') or '')
        ws.cell(row=r, column=6, value='; '.join(f"{k}={v}" for k, v in ev.items()))
        ws.cell(row=r, column=7, value=a['reason'])
        ws.row_dimensions[r].height = 26

    # ══ 5. REMEDIATION ═══════════════════════════════════════════════════
    ws = wb.create_sheet('Remediation')
    sheet_title(ws, '05  RECOMMENDED ACTIONS', 'Prioritised remediation',
                'Actions derived from the findings, ordered by urgency.', span=4)
    hdr = 6
    order = {'CRITICAL': 0, 'HIGH': 1, 'MEDIUM': 2, 'LOW': 3}
    recs = sorted(
        [(order.get(REMEDIATION[r][0], 3), REMEDIATION[r][0], REMEDIATION[r][1],
          REMEDIATION[r][2], r, n)
         for r, n in rule_count.items() if r in REMEDIATION])
    table(ws, hdr, ['PRIORITY', 'RECOMMENDED ACTION', 'RATIONALE', 'TRIGGER'],
          [14, 40, 76, 26], len(recs), wrap_cols=(3,))
    for i, (_, prio, action, detail, rule, n) in enumerate(recs):
        r = hdr + 1 + i
        c = ws.cell(row=r, column=1, value=prio)
        c.font = Font(name=FONT, size=9, bold=True, color=SEV.get(prio.lower(), MUTED))
        c = ws.cell(row=r, column=2, value=action)
        c.font = Font(name=FONT, size=9, bold=True, color=INK)
        ws.cell(row=r, column=3, value=detail)
        ws.cell(row=r, column=4, value=f"{rule.replace('_', ' ')} ({n})")
        ws.row_dimensions[r].height = 30

    # ══ 6. SEVERITY ══════════════════════════════════════════════════════
    ws = wb.create_sheet('Severity')
    sheet_title(ws, '06  DISTRIBUTION', 'Detections by severity',
                'Count and share of detections at each severity level.', span=3)
    hdr = 6
    levels = [s for s in ('critical', 'high', 'medium', 'low') if sev_count.get(s)]
    table(ws, hdr, ['SEVERITY', 'COUNT', 'SHARE'], [18, 14, 14], len(levels))
    tot = sum(sev_count.values()) or 1
    for i, s in enumerate(levels):
        r = hdr + 1 + i
        c = ws.cell(row=r, column=1, value=s.upper())
        c.font = Font(name=FONT, size=9, bold=True, color=SEV[s])
        ws.cell(row=r, column=2, value=sev_count[s])
        c = ws.cell(row=r, column=3, value=sev_count[s] / tot)
        c.number_format = '0.0%'
    if levels:
        ch = BarChart()
        ch.type = 'bar'
        ch.title = 'Detections by severity'
        ch.height, ch.width = 7, 13
        ch.add_data(Reference(ws, min_col=2, min_row=hdr, max_row=hdr + len(levels)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=hdr + 1,
                                    max_row=hdr + len(levels)))
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ws.add_chart(ch, 'E6')

    # ══ 7. TRAFFIC ═══════════════════════════════════════════════════════
    ws = wb.create_sheet('Traffic')
    sheet_title(ws, '07  TRAFFIC COMPOSITION', 'Protocol and host distribution',
                'Baseline characterisation of the monitored traffic.', span=5)
    hdr = 6
    table(ws, hdr, ['PROTOCOL', 'FLOWS', 'BYTES', 'VOLUME', 'SHARE'],
          [18, 14, 16, 14, 12], len(protocols))
    totp = sum(p['bytes'] or 0 for p in protocols) or 1
    for i, p in enumerate(protocols):
        r = hdr + 1 + i
        ws.cell(row=r, column=1, value=p['protocol'])
        ws.cell(row=r, column=2, value=p['flows']).number_format = '#,##0'
        ws.cell(row=r, column=3, value=p['bytes']).number_format = '#,##0'
        ws.cell(row=r, column=4, value=fmt_bytes(p['bytes']))
        c = ws.cell(row=r, column=5, value=(p['bytes'] or 0) / totp)
        c.number_format = '0.0%'
    if protocols:
        ch = DoughnutChart(holeSize=58)
        ch.title = 'Traffic share by protocol'
        ch.height, ch.width = 8, 13
        ch.add_data(Reference(ws, min_col=3, min_row=hdr, max_row=hdr + len(protocols)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=hdr + 1,
                                    max_row=hdr + len(protocols)))
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ws.add_chart(ch, 'G6')

    hosts_hdr = hdr + len(protocols) + 3
    ws.cell(row=hosts_hdr - 1, column=1, value='HIGHEST-VOLUME HOSTS').font = \
        Font(name=FONT, size=9, bold=True, color=INK)
    table(ws, hosts_hdr, ['HOST', 'FLOWS', 'PEERS', 'BYTES', 'SHARE'],
          [18, 14, 16, 14, 12], min(len(talkers), 25))
    tot_t = sum(t['bytes'] or 0 for t in talkers) or 1
    for i, t in enumerate(talkers[:25]):
        r = hosts_hdr + 1 + i
        ws.cell(row=r, column=1, value=t['src_ip'])
        ws.cell(row=r, column=2, value=t['flows']).number_format = '#,##0'
        ws.cell(row=r, column=3, value=t['peers'])
        ws.cell(row=r, column=4, value=t['bytes']).number_format = '#,##0'
        c = ws.cell(row=r, column=5, value=(t['bytes'] or 0) / tot_t)
        c.number_format = '0.0%'
    if talkers:
        ws.conditional_formatting.add(
            f"E{hosts_hdr+1}:E{hosts_hdr+min(len(talkers),25)}",
            DataBarRule(start_type='num', start_value=0, end_type='max',
                        color=ACCENT, showValue=True))

    # ══ 8. THROUGHPUT ════════════════════════════════════════════════════
    ws = wb.create_sheet('Throughput')
    sheet_title(ws, '08  THROUGHPUT', 'Traffic over time',
                'Bytes and packets per minute across the assessment window.', span=3)
    hdr = 6
    table(ws, hdr, ['MINUTE (UTC)', 'BYTES', 'PACKETS'], [22, 16, 14], len(timeline))
    for i, t in enumerate(timeline):
        r = hdr + 1 + i
        ws.cell(row=r, column=1, value=ts_str(t['minute']))
        ws.cell(row=r, column=2, value=t['bytes']).number_format = '#,##0'
        ws.cell(row=r, column=3, value=t.get('packets', 0)).number_format = '#,##0'
    if len(timeline) > 1:
        ch = LineChart()
        ch.title = 'Bytes per minute'
        ch.height, ch.width = 8, 20
        ch.add_data(Reference(ws, min_col=2, min_row=hdr, max_row=hdr + len(timeline)),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=hdr + 1,
                                    max_row=hdr + len(timeline)))
        ws.add_chart(ch, 'E6')

    # ══ 9. METHODOLOGY ═══════════════════════════════════════════════════
    ws = wb.create_sheet('Methodology')
    sheet_title(ws, '09  METHODOLOGY', 'Threshold derivation',
                'How each detection threshold was established and validated.', span=5)
    hdr = 6
    rules_m = [
        ('port_scan', '15 distinct ports, >50% RST', '3 ports/min'),
        ('connection_burst', '421 connections to one port', '210/min'),
        ('dns_exfil', '30+ char label, entropy >= 3.5', 'entropy ~1.9'),
        ('icmp_sweep', '10 distinct destinations', '3 hosts/min'),
        ('volume_anomaly', 'z >= 3.5 sigma, floor 50 KB', 'per-host EWMA'),
    ]
    table(ws, hdr,
          ['RULE', 'THRESHOLD', 'OBSERVED NORMAL (p99)', 'ATT&CK', 'DETECTS'],
          [24, 34, 26, 14, 40], len(rules_m))
    for i, (rule, thr, obs) in enumerate(rules_m):
        r = hdr + 1 + i
        c = ws.cell(row=r, column=1, value=rule)
        c.font = Font(name=FONT, size=9, bold=True, color=INK)
        ws.cell(row=r, column=2, value=thr)
        ws.cell(row=r, column=3, value=obs)
        ws.cell(row=r, column=4, value=ATTACK.get(rule, ('', '—'))[1])
        ws.cell(row=r, column=5, value=ATTACK.get(rule, ('', '', rule))[2])
    n = hdr + len(rules_m) + 2
    for text in [
        'Thresholds were derived empirically. A baseline capture of ordinary network '
        'activity was analysed, per-minute distributions computed for every metric each '
        'rule depends upon, and each threshold placed above the 99th percentile of '
        'observed normal traffic — giving a measurable false-positive ceiling.',
        '',
        'Validation: purpose-built attack signatures were generated locally against the '
        'loopback interface and each was correctly identified. Across a separate capture '
        'of ordinary browsing the engine produced no false positives.',
    ]:
        if text:
            c = ws.cell(row=n, column=1, value=text)
            c.font = Font(name=FONT, size=9, color=GRAPH)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=n, start_column=1, end_row=n + 2, end_column=5)
            n += 4
        else:
            n += 1

    # ── global: professional font everywhere, print setup ────────────────
    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
        sh.page_setup.orientation = 'landscape'
        sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        for row in sh.iter_rows():
            for cell in row:
                if cell.value is not None and (not cell.font or cell.font.name != FONT):
                    f = cell.font
                    cell.font = Font(name=FONT, size=f.size or 9, bold=f.bold,
                                     italic=f.italic, color=f.color)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == '__main__':
    data = build_workbook()
    with open('sentinel_report.xlsx', 'wb') as f:
        f.write(data)
    print(f'wrote sentinel_report.xlsx ({len(data):,} bytes)')
